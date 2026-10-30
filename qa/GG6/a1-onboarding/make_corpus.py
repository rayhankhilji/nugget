#!/usr/bin/env python3
"""Write a synthetic "new user's personal files" corpus (about 150 files) for the A1 QA scenarios.

Everything here is invented: the owner "Sam Tellwright", their contacts and companies are fictional.
Usage: python3 -I make_corpus.py <out-dir>
Writes <out-dir>/corpus/... and <out-dir>/questions.json (10 day-one questions with the file
that answers each one). Deterministic (seeded).
"""
import json
import os
import random
import sys
import zlib
from datetime import date, timedelta

OUT = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else "a1-corpus")
ROOT = os.path.join(OUT, "corpus")
rnd = random.Random(42)
ME = "Sam Tellwright"
MY_MAIL = "sam@tellwright.test"


def write(rel, text, mode="w"):
    p = os.path.join(ROOT, rel)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, mode, encoding=None if "b" in mode else "utf-8") as f:
        f.write(text)
    return rel


def eml(rel, frm, subject, day, body, to=MY_MAIL):
    return write(
        rel,
        f"From: {frm}\nTo: {to}\nSubject: {subject}\nDate: {day.strftime('%a, %d %b %Y')} 10:12:00 +0000\n"
        f"Message-ID: <{zlib.crc32((rel + subject).encode())}@mail.test>\nMIME-Version: 1.0\n"
        f"Content-Type: text/plain; charset=utf-8\n\n{body}\n",
    )


def ics(rel, summary, start, minutes, location="", desc=""):
    end_h, end_m = divmod(start[1] * 60 + start[2] + minutes, 60)
    d = start[0].strftime("%Y%m%d")
    return write(
        rel,
        "BEGIN:VCALENDAR\r\nVERSION:2.0\r\nPRODID:-//a1-qa//EN\r\nBEGIN:VEVENT\r\n"
        f"UID:{zlib.crc32(rel.encode())}@a1.test\r\nDTSTAMP:20261001T080000Z\r\n"
        f"DTSTART:{d}T{start[1]:02d}{start[2]:02d}00\r\nDTEND:{d}T{end_h:02d}{end_m:02d}00\r\n"
        f"SUMMARY:{summary}\r\nLOCATION:{location}\r\nDESCRIPTION:{desc}\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n",
    )


questions = []


def q(question, file, must):
    questions.append({"q": question, "file": file, "must": must})


# ---- the ten answers (needles) ----
q(
    "when is my dentist appointment",
    ics("calendar/2026-11-12-dentist.ics", "Dentist - Dr. Okafor (cleaning)", (date(2026, 11, 12), 9, 30), 45,
        "Brightsmile Dental, 14 Larch Street", "Six-month cleaning and checkup. Bring the new insurance card."),
    "Okafor",
)
q(
    "what did Priya say about the budget",
    eml("mail/2026-10-02-priya-budget.eml", "Priya Raman <priya@fernhollow.test>", "Re: Q4 budget for the community garden",
        date(2026, 10, 2),
        "Hi Sam,\n\nI went through the Q4 numbers. We're about 600 over on tools, so I'd cut the second\n"
        "rain barrel and push the greenhouse panels to January. Everything else in the budget looks fine.\n"
        "Can you send the revised sheet to Tomasz by Friday?\n\nPriya"),
    "rain barrel",
)
q(
    "what is the wifi password at the cabin",
    write("notes/cabin-at-lake-mirren.md",
          "# Cabin at Lake Mirren\n\nKey is in the lockbox by the shed, code 4417.\n\n"
          "WiFi network: MirrenLoft, password: otter-canoe-1988\n\nBring firewood starter; the stove pulls badly when cold.\n"),
    "otter-canoe-1988",
)
q(
    "when does my car insurance renew",
    "documents/insurance/auto-policy-2026.pdf",
    "March 3, 2027",
)
q(
    "gift ideas for mom's birthday",
    write("notes/gifts/mom-birthday-ideas.txt",
          "Mom birthday (Dec 9) ideas:\n- the blue enamel teapot from the market on Hollis Ave\n- pottery class voucher\n"
          "- reframe the 1994 beach photo\nBudget about 80.\n"),
    "teapot",
)
q(
    "flight to Lisbon confirmation number",
    eml("mail/2026-09-18-gullwing-booking.eml", "Gullwing Air <bookings@gullwing-air.test>",
        "Your booking is confirmed: Boston to Lisbon", date(2026, 9, 18),
        "Dear Sam Tellwright,\n\nThank you for booking with Gullwing Air.\nConfirmation number: QX7Z4M\n"
        "Flight GW 214, Boston (BOS) to Lisbon (LIS), departs 14 Nov 2026 at 21:40.\nSeat 23A. One checked bag included.\n"),
    "QX7Z4M",
)
q(
    "banana bread recipe",
    write("notes/recipes/banana-bread.md",
          "# Grandma Ilse's banana bread\n\n- 3 very ripe bananas\n- 75 g melted butter\n- 150 g sugar\n- 1 egg\n"
          "- 1 tsp baking soda, pinch of salt\n- 190 g flour\n\nBake 60 min at 175 C. Walnuts optional.\n"),
    "ripe bananas",
)
q(
    "what's the plumber's phone number",
    write("notes/house/contractors.txt",
          "House contacts\n\nPlumber: Dario Venn, Venn & Sons Plumbing, 555-0143 (fixed the kitchen drain in May)\n"
          "Electrician: Hollis Mbeki, 555-0178\nRoofer: ask Priya, she used someone good\n"),
    "555-0143",
)
q(
    "when does my passport expire",
    "documents/travel/travel-documents.docx",
    "2029",
)
q(
    "notes from my 1:1 with Marcus",
    write("notes/work/1on1-marcus-2026-10-20.md",
          "# 1:1 with Marcus Ebbe, 2026-10-20\n\n- Marcus wants the onboarding survey shipped before the offsite\n"
          "- I asked for a second reviewer on the pricing page; he'll ask Lena\n- Career: talk about the lead role again in January\n"),
    "offsite",
)

# PDF and DOCX needles (extractor supports both).
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    p = os.path.join(ROOT, "documents/insurance/auto-policy-2026.pdf")
    os.makedirs(os.path.dirname(p), exist_ok=True)
    c = canvas.Canvas(p, pagesize=letter)
    y = 720
    for line in ["Harborline Mutual - Auto Policy Declarations", f"Policyholder: {ME}",
                 "Policy number HM-55-201934", "Vehicle: 2019 hatchback, blue",
                 "Policy period: March 3, 2026 to March 3, 2027", "This policy renews on March 3, 2027.",
                 "Annual premium: 1,184.00"]:
        c.drawString(72, y, line)
        y -= 20
    c.save()
except ImportError:
    print("reportlab missing: no PDF", file=sys.stderr)
try:
    import docx

    d = docx.Document()
    d.add_heading("Travel documents", 1)
    d.add_paragraph(f"Passport: {ME}, number X4410925, issued 2019-06-11, expires 2029-06-10.")
    d.add_paragraph("Global entry renewal due 2027. Copies of both are in the fire safe.")
    p = os.path.join(ROOT, "documents/travel/travel-documents.docx")
    os.makedirs(os.path.dirname(p), exist_ok=True)
    d.save(p)
except ImportError:
    print("python-docx missing: no DOCX", file=sys.stderr)

# ---- filler: realistic distractors ----
people = ["Priya Raman", "Tomasz Kral", "Lena Osei", "Marcus Ebbe", "Ines Varga", "Hugo Pell", "Mina Sato", "Dev Arora"]
topics = ["garden", "budget", "offsite", "pricing page", "school pickup", "book club", "running", "taxes", "car", "kitchen"]
start = date(2026, 6, 1)
n = 0
for i in range(50):  # journal entries
    d = start + timedelta(days=i * 3)
    t1, t2 = rnd.sample(topics, 2)
    who = rnd.choice(people)
    write(f"journal/{d.isoformat()}.md",
          f"# {d.strftime('%A %d %B %Y')}\n\nTalked with {who} about the {t1}. Slept badly. "
          f"Need to sort out the {t2} this week.\n\n- [ ] call about the {t2}\n- [x] walk 5 km\n")
    n += 1
for i in range(35):  # mail
    d = start + timedelta(days=i * 4 + 1)
    who = rnd.choice(people)
    topic = rnd.choice(topics)
    first = who.split()[0].lower()
    eml(f"mail/{d.isoformat()}-{first}-{i}.eml", f"{who} <{first}@fernhollow.test>", f"About the {topic}", d,
        f"Hi Sam,\n\nQuick note on the {topic}: let's look at it again next week. "
        f"I think we're mostly fine but want a second look.\n\n{who.split()[0]}")
    n += 1
appts = [("Eye doctor - Dr. Lund", "Clearview Optics"), ("Vet - Biscuit vaccines", "Paws on Pine"),
         ("Haircut", "Shear Luck"), ("Book club", "Library room 2"), ("Parent-teacher meeting", "Rowan Elementary"),
         ("Car service", "Kessler Auto"), ("Physio", "Movewell Clinic"), ("Team offsite planning", "Room 4B")]
for i in range(25):  # calendar
    d = date(2026, 10, 1) + timedelta(days=i * 3)
    s, loc = appts[i % len(appts)]
    ics(f"calendar/{d.isoformat()}-{i}.ics", s, (d, 8 + i % 9, 0), 60, loc)
    n += 1
for i in range(20):  # misc notes (txt)
    t = rnd.choice(topics)
    write(f"notes/misc/note-{i:02d}.txt",
          f"Random thought about {t}: maybe ask {rnd.choice(people)}. Price was around {rnd.randint(20, 900)}.\n")
    n += 1
write("notes/house/budget-2026.csv", "category,planned,actual\ngroceries,600,642\nutilities,210,198\ncar,150,320\n")

# A few files the default extensions skip, so the folder looks like a real one.
write("downloads/setup-tool.exe", "MZ" + "\0" * 64)
write("downloads/photos-backup.zip", "PK\x03\x04" + "\0" * 64)

with open(os.path.join(OUT, "questions.json"), "w") as f:
    json.dump(questions, f, indent=1)
count = sum(len(fs) for _, _, fs in os.walk(ROOT))
print(json.dumps({"root": ROOT, "files": count, "questions": len(questions)}))
